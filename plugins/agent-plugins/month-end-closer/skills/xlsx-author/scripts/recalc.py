#!/usr/bin/env python3

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

EXCEL_ERRORS = {
    "#NULL!",
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#NUM!",
    "#N/A",
    "#SPILL!",
    "#CALC!",
}


def find_libreoffice():
    """Return the LibreOffice executable path if available."""

    for command in ("libreoffice", "soffice"):
        path = shutil.which(command)
        if path:
            return path

    # Common macOS installation location.
    macos_path = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if macos_path.exists():
        return str(macos_path)

    return None


def recalculate_workbook(workbook_path, timeout):
    office = find_libreoffice()

    if not office:
        raise RuntimeError(
            "LibreOffice was not found. Install LibreOffice or make "
            "'libreoffice'/'soffice' available on PATH."
        )

    workbook_path = workbook_path.resolve()

    with tempfile.TemporaryDirectory(prefix="recalc-") as temp_dir:
        temp_dir = Path(temp_dir)
        output_dir = temp_dir / "output"
        profile_dir = temp_dir / "profile"

        output_dir.mkdir()
        profile_dir.mkdir()

        command = [
            office,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_dir),
            str(workbook_path),
        ]

        try:
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=timeout,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(
                f"LibreOffice recalculation timed out after {timeout} seconds."
            ) from exc

        if result.returncode != 0:
            message = result.stderr.strip() or result.stdout.strip()
            raise RuntimeError(f"LibreOffice failed to recalculate workbook: {message}")

        recalculated_path = output_dir / workbook_path.name

        if not recalculated_path.exists():
            raise RuntimeError(
                "LibreOffice completed without producing a recalculated workbook."
            )

        shutil.copy2(recalculated_path, workbook_path)


def inspect_formula_errors(workbook_path):
    formula_workbook = load_workbook(workbook_path, data_only=False)
    value_workbook = load_workbook(workbook_path, data_only=True)

    total_formulas = 0
    errors = {}

    for sheet_name in formula_workbook.sheetnames:
        formula_sheet = formula_workbook[sheet_name]
        value_sheet = value_workbook[sheet_name]

        for row in formula_sheet.iter_rows():
            for formula_cell in row:
                if formula_cell.data_type != "f":
                    continue

                total_formulas += 1

                value_cell = value_sheet[formula_cell.coordinate]
                value = value_cell.value

                if value_cell.data_type == "e" or value in EXCEL_ERRORS:
                    error = str(value)
                    location = f"{sheet_name}!{formula_cell.coordinate}"

                    if error not in errors:
                        errors[error] = {
                            "count": 0,
                            "locations": [],
                        }

                    errors[error]["count"] += 1
                    errors[error]["locations"].append(location)

    total_errors = sum(details["count"] for details in errors.values())

    return total_formulas, total_errors, errors


def main():
    if len(sys.argv) not in (2, 3):
        print(
            json.dumps(
                {
                    "status": "error",
                    "message": "Usage: recalc.py <workbook.xlsx> [timeout_seconds]",
                },
                indent=2,
            )
        )
        return 2

    workbook_path = Path(sys.argv[1])

    try:
        timeout = int(sys.argv[2]) if len(sys.argv) == 3 else 30
    except ValueError:
        print(
            json.dumps(
                {
                    "status": "error",
                    "message": "timeout_seconds must be an integer.",
                },
                indent=2,
            )
        )
        return 2

    if not workbook_path.exists():
        print(
            json.dumps(
                {
                    "status": "error",
                    "message": f"Workbook not found: {workbook_path}",
                },
                indent=2,
            )
        )
        return 2

    if workbook_path.suffix.lower() != ".xlsx":
        print(
            json.dumps(
                {
                    "status": "error",
                    "message": "Only .xlsx workbooks are supported.",
                },
                indent=2,
            )
        )
        return 2

    try:
        recalculate_workbook(workbook_path, timeout)

        total_formulas, total_errors, errors = inspect_formula_errors(workbook_path)

    except Exception as exc:
        print(
            json.dumps(
                {
                    "status": "error",
                    "message": str(exc),
                },
                indent=2,
            )
        )
        return 2

    result = {
        "status": "success" if total_errors == 0 else "errors_found",
        "total_errors": total_errors,
        "total_formulas": total_formulas,
        "error_summary": errors,
    }

    print(json.dumps(result, indent=2))

    return 0 if total_errors == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
